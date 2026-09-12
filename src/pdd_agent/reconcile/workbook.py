"""Read a registered ACM0022 ER calculation workbook with cell lineage.

Every number carries the sheet and A1 cell it was read from, so the Inegol
config re-sourcing and the reconcile report can cite exact workbook cells.
Columns are located by header cell text, never by fixed letters (RISK-05-02):
headers are matched with ``str(cell).strip().startswith(prefix)`` after
collapsing whitespace.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from pdd_agent.calc.constants import DECAY_RATE_BY_CLIMATE_ZONE

#: Default location of the registered Inegol ER workbook (ASM-001, git-ignored).
INEGOL_ER_WORKBOOK_DEFAULT = Path(
    "data/runs/tinh-20260827/ws/PDD_Portable_Workspace_20260827/inputs/fresh_registry/"
    "verra_public_download_3908_20260824/project_3908/extracted/"
    "ER Calculation.v04_21.07.2025.xlsx"
)

#: Environment variable overriding the workbook location (ASM-001).
INEGOL_ER_WORKBOOK_ENV_VAR = "PDD_INEGOL_ER_WORKBOOK"

#: ASM-002: the ECBL cross-check constants. The registered workbook's *Project
#: Emissions* sheet gives EFgrid = 0.541 tCO2/MWh (cell B11) and TDL = 0.09101
#: (cell B12); the separated baseline must equal MWh x 0.541 x 1.09101.
ECBL_GRID_EF = 0.541
ECBL_TDL_MULTIPLIER = 1.09101
ECBL_CROSS_CHECK_TOLERANCE = 1.0  # tCO2e per year

#: Rows whose Days cell is below this are pre-crediting stubs (ASM-007).
STUB_DAYS_THRESHOLD = 300

#: Crediting-year labels: 01/01/2021 through 01/01/2027 starts.
_YEAR_LABEL_RE = re.compile(r"^01/01/202([1-7])\b")

#: Workbook row-label prefix (normalised) -> engine waste-type key.
_LABEL_PREFIX_TO_WASTE_TYPE: tuple[tuple[str, str], ...] = (
    ("wood", "wood"),
    ("pulp, paper", "paper_cardboard"),
    ("food", "food_waste"),
    ("textiles", "textiles"),
    ("garden", "garden_waste"),
    ("glass, plastic, metal", "inert"),
)


def inegol_workbook_path() -> Path:
    """Resolve the registered Inegol ER workbook path (ASM-001)."""
    override = os.environ.get(INEGOL_ER_WORKBOOK_ENV_VAR)
    if override:
        return Path(override)
    return INEGOL_ER_WORKBOOK_DEFAULT


def infer_climate_zone(decay_rates: dict[str, float]) -> str | None:
    """Infer the IPCC climate zone from registered decay rates (S-5).

    Returns the single zone in ``DECAY_RATE_BY_CLIMATE_ZONE`` whose table
    matches **all** supplied values within 1e-9, and ``None`` when zero or
    more than one zone matches.
    """
    matches = [
        zone
        for zone, table in DECAY_RATE_BY_CLIMATE_ZONE.items()
        if all(
            waste_type in table and abs(table[waste_type] - value) <= 1e-9
            for waste_type, value in decay_rates.items()
        )
    ]
    if len(matches) == 1:
        return matches[0]
    return None


@dataclass(frozen=True)
class LineageValue:
    """One number and where it came from (cell in A1 notation)."""

    value: float
    sheet: str
    cell: str


@dataclass
class RegisteredSchedule:
    """Crediting years only; stub rows removed."""

    year_labels: list[str] = field(default_factory=list)
    baseline: list[LineageValue] = field(default_factory=list)
    project: list[LineageValue] = field(default_factory=list)
    leakage: list[LineageValue] = field(default_factory=list)
    net: list[LineageValue] = field(default_factory=list)


@dataclass
class RegisteredAcm0022Workbook:
    """Everything the engine and the reconcile report need."""

    path: Path
    schedule: RegisteredSchedule = field(default_factory=RegisteredSchedule)
    decay_rates: dict[str, LineageValue] = field(default_factory=dict)
    doc_fractions: dict[str, LineageValue] = field(default_factory=dict)
    composition: dict[str, LineageValue] = field(default_factory=dict)
    waste_by_year: list[LineageValue] = field(default_factory=list)
    electricity_mwh_by_year: list[LineageValue] = field(default_factory=list)
    biomethanization_fraction: LineageValue | None = None
    separated_baseline_by_year: list[LineageValue] = field(default_factory=list)


def _norm(text: Any) -> str:
    return " ".join(str(text).split()).lower()


def _map_label_to_waste_type(label: str) -> str | None:
    lowered = _norm(label)
    for prefix, waste_type in _LABEL_PREFIX_TO_WASTE_TYPE:
        if lowered.startswith(prefix):
            return waste_type
    return None


def _get_sheet(wb: Any, name: str) -> Any:
    if name not in wb.sheetnames:
        raise ValueError(
            f'registered workbook is missing required sheet "{name}"; found sheets: {wb.sheetnames}'
        )
    return wb[name]


def _rows(ws: Any) -> list[list[Any]]:
    return list(ws.iter_rows(values_only=False))


def _find_header_row(
    ws: Any, sheet: str, required: list[str], startswith: list[str] | None = None
) -> tuple[int, list[Any]]:
    """Find the first row containing every ``required`` exact header.

    ``startswith`` prefixes must each match at least one cell in the same row.
    Returns the 0-based row index and the row's cells.
    """
    startswith = startswith or []
    for idx, row in enumerate(_rows(ws)):
        texts = [_norm(c.value) for c in row if c.value is not None]
        if all(any(t == r for t in texts) for r in required) and all(
            any(t.startswith(p) for t in texts) for p in startswith
        ):
            return idx, row
    want = required + [p + "*" for p in startswith]
    raise ValueError(f'sheet "{sheet}": header row with {want} not found')


def _column_matching(cells: list[Any], prefix: str) -> int:
    """0-based column index of the first cell starting with ``prefix``."""
    for col, cell in enumerate(cells):
        if cell.value is not None and _norm(cell.value).startswith(prefix):
            return col
    raise ValueError(f'header cell starting with "{prefix}" not found')


def _read_schedule(ws: Any, sheet: str) -> RegisteredSchedule:
    header_idx, header = _find_header_row(
        ws,
        sheet,
        required=["days", "year"],
        startswith=["baseline emissions", "project emissions", "emission reductions"],
    )
    texts = [_norm(c.value) if c.value is not None else "" for c in header]
    days_col = texts.index("days")
    year_col = texts.index("year")
    baseline_col = next(i for i, t in enumerate(texts) if t.startswith("baseline emissions"))
    project_col = next(i for i, t in enumerate(texts) if t.startswith("project emissions"))
    leakage_col = next(i for i, t in enumerate(texts) if "leak" in t)
    net_col = next(i for i, t in enumerate(texts) if t.startswith("emission reductions"))

    schedule = RegisteredSchedule()
    for row in _rows(ws)[header_idx + 1 :]:
        first = _norm(row[0].value) if row and row[0].value is not None else ""
        if first == "total":
            break
        days = row[days_col].value if days_col < len(row) else None
        if days is None or float(days) < STUB_DAYS_THRESHOLD:
            continue  # one-day pre-crediting stub row (ASM-007)

        def _at(col: int) -> LineageValue:
            cell = row[col]
            return LineageValue(value=float(cell.value), sheet=sheet, cell=cell.coordinate)

        schedule.year_labels.append(str(row[year_col].value))
        schedule.baseline.append(_at(baseline_col))
        schedule.project.append(_at(project_col))
        schedule.leakage.append(_at(leakage_col))
        schedule.net.append(_at(net_col))
    if not schedule.year_labels:
        raise ValueError(f'sheet "{sheet}": no crediting rows found under the "Days"/"Year" header')
    return schedule


def _read_decay_and_doc(
    ws: Any, sheet: str
) -> tuple[dict[str, LineageValue], dict[str, LineageValue]]:
    header_idx, header = _find_header_row(ws, sheet, required=["docj", "kj"])
    texts = [_norm(c.value) if c.value is not None else "" for c in header]
    doc_col = texts.index("docj")
    k_col = texts.index("kj")
    decay_rates: dict[str, LineageValue] = {}
    doc_fractions: dict[str, LineageValue] = {}
    for row in _rows(ws)[header_idx + 1 :]:
        label = next(
            (str(c.value) for c in row if isinstance(c.value, str) and c.value.strip()),
            "",
        )
        if not label:
            continue
        waste_type = _map_label_to_waste_type(label)
        if waste_type is None or waste_type == "inert":
            continue
        doc_cell = row[doc_col] if doc_col < len(row) else None
        k_cell = row[k_col] if k_col < len(row) else None
        if doc_cell is None or doc_cell.value is None or k_cell is None or k_cell.value is None:
            continue
        # The workbook writes DOC as a percentage (43 = 0.43); k is a fraction.
        doc_fractions[waste_type] = LineageValue(
            value=float(doc_cell.value) / 100.0, sheet=sheet, cell=doc_cell.coordinate
        )
        decay_rates[waste_type] = LineageValue(
            value=float(k_cell.value), sheet=sheet, cell=k_cell.coordinate
        )
    if not decay_rates:
        raise ValueError(f'sheet "{sheet}": no DOCj/kj parameter rows found')
    return decay_rates, doc_fractions


def _read_waste_projection(
    ws: Any, sheet: str
) -> tuple[dict[str, LineageValue], list[LineageValue], LineageValue | None]:
    rows = _rows(ws)
    # Composition table: header row with "Waste Type" + "At the Integrated Solid...".
    comp_header_idx: int | None = None
    for idx, row in enumerate(rows):
        texts = [_norm(c.value) if c.value is not None else "" for c in row]
        if "waste type" in texts and any(t.startswith("at the integrated solid") for t in texts):
            comp_header_idx = idx
            comp_header = row
            break
    if comp_header_idx is None:
        raise ValueError(
            f'sheet "{sheet}": composition header with "Waste Type" and '
            '"At the Integrated Solid..." not found'
        )
    comp_texts = [_norm(c.value) if c.value is not None else "" for c in comp_header]
    label_col = comp_texts.index("waste type")
    comp_col = next(i for i, t in enumerate(comp_texts) if t.startswith("at the integrated solid"))
    composition: dict[str, LineageValue] = {}
    for row in rows[comp_header_idx + 1 :]:
        if label_col >= len(row) or not isinstance(row[label_col].value, str):
            continue
        label = row[label_col].value
        if not label.strip():
            continue
        if _norm(label) == "waste type":
            break  # a second composition table follows; keep the first
        waste_type = _map_label_to_waste_type(label)
        if waste_type is None or comp_col >= len(row) or row[comp_col].value is None:
            continue
        composition[waste_type] = LineageValue(
            value=float(row[comp_col].value), sheet=sheet, cell=row[comp_col].coordinate
        )
    if not composition:
        raise ValueError(f'sheet "{sheet}": no "At the Integrated Solid..." composition rows found')

    # Waste-by-year table: header row with a "Total Waste..." column.
    year_header_idx: int | None = None
    for idx, row in enumerate(rows):
        texts = [_norm(c.value) if c.value is not None else "" for c in row]
        if any(t.startswith("total waste") for t in texts):
            year_header_idx = idx
            year_header = row
            break
    if year_header_idx is None:
        raise ValueError(f'sheet "{sheet}": column starting with "Total Waste" not found')
    year_texts = [_norm(c.value) if c.value is not None else "" for c in year_header]
    total_col = next(i for i, t in enumerate(year_texts) if t.startswith("total waste"))
    biometh_col = next(
        (i for i, t in enumerate(year_texts) if t.startswith("biomethanization")), None
    )
    year_col = next((i for i, t in enumerate(year_texts) if t == "year"), 0)
    waste_by_year: list[LineageValue] = []
    biomethanization_fraction: LineageValue | None = None
    for row in rows[year_header_idx + 1 :]:
        if year_col >= len(row) or row[year_col].value is None:
            continue
        label = str(row[year_col].value)
        if not _YEAR_LABEL_RE.match(label):
            continue
        total_cell = row[total_col]
        waste_by_year.append(
            LineageValue(value=float(total_cell.value), sheet=sheet, cell=total_cell.coordinate)
        )
        if biomethanization_fraction is None and biometh_col is not None:
            bio_cell = row[biometh_col]
            if bio_cell.value is not None:
                biomethanization_fraction = LineageValue(
                    value=float(bio_cell.value) / float(total_cell.value),
                    sheet=sheet,
                    cell=bio_cell.coordinate,
                )
    if not waste_by_year:
        raise ValueError(f'sheet "{sheet}": no "Total Waste" rows for 2021-2027 found')
    return composition, waste_by_year, biomethanization_fraction


def _read_electricity_with_cross_check(
    wb: Any,
) -> tuple[list[LineageValue], list[LineageValue]]:
    """Read the ECBL series and validate it against the separated baseline (ASM-002)."""
    project_sheet = "Project Emissions"
    baseline_sheet = "Baseline Emissions (Total)"
    ws = _get_sheet(wb, project_sheet)
    header_idx, header = _find_header_row(
        ws, sheet=project_sheet, required=["year"], startswith=["ecbl"]
    )
    texts = [_norm(c.value) if c.value is not None else "" for c in header]
    year_col = texts.index("year")
    ecbl_col = next(i for i, t in enumerate(texts) if t.startswith("ecbl"))
    electricity: list[LineageValue] = []
    electricity_years: list[str] = []
    for row in _rows(ws)[header_idx + 1 :]:
        # The sheet holds several year-labelled tables (ECBL, flare, AD);
        # the ECBL series is the first contiguous run of crediting-year rows.
        if year_col >= len(row) or row[year_col].value is None:
            if electricity:
                break
            continue
        label = str(row[year_col].value)
        if not _YEAR_LABEL_RE.match(label):
            if electricity:
                break
            continue
        cell = row[ecbl_col]
        electricity.append(
            LineageValue(value=float(cell.value), sheet=project_sheet, cell=cell.coordinate)
        )
        electricity_years.append(label)
    if not electricity:
        raise ValueError(f'sheet "{project_sheet}": no "ECBL" rows for 2021-2027 found')

    bws = _get_sheet(wb, baseline_sheet)
    b_header_idx, b_header = _find_header_row(
        bws, sheet=baseline_sheet, required=[], startswith=["baseline emissions from separate"]
    )
    b_texts = [_norm(c.value) if c.value is not None else "" for c in b_header]
    sep_col = next(
        i for i, t in enumerate(b_texts) if t.startswith("baseline emissions from separate")
    )
    b_year_col = next((i for i, t in enumerate(b_texts) if t == "crediting period"), 0)
    separated: list[LineageValue] = []
    separated_years: list[str] = []
    for row in _rows(bws)[b_header_idx + 1 :]:
        if b_year_col >= len(row) or row[b_year_col].value is None:
            if separated:
                break
            continue
        label = str(row[b_year_col].value)
        if not _YEAR_LABEL_RE.match(label):
            if separated:
                break
            continue
        cell = row[sep_col]
        separated.append(
            LineageValue(value=float(cell.value), sheet=baseline_sheet, cell=cell.coordinate)
        )
        separated_years.append(label)
    if len(separated) != len(electricity):
        raise ValueError(
            f'"ECBL" cross-check: {len(electricity)} electricity years vs '
            f"{len(separated)} separated-baseline years"
        )
    for year_label, mwh, sep in zip(electricity_years, electricity, separated):
        expected = mwh.value * ECBL_GRID_EF * ECBL_TDL_MULTIPLIER
        if abs(expected - sep.value) > ECBL_CROSS_CHECK_TOLERANCE:
            raise ValueError(
                f'"ECBL" cross-check failed at {mwh.cell} ({year_label}): {mwh.value} MWh x '
                f"{ECBL_GRID_EF} x {ECBL_TDL_MULTIPLIER} = {expected:.1f} tCO2e, "
                f'but "{baseline_sheet}" {sep.cell} holds {sep.value}'
            )
    return electricity, separated


def read_acm0022_workbook(path: Path | str) -> RegisteredAcm0022Workbook:
    """Parse a registered ACM0022 ER calculation workbook with cell lineage.

    Raises ``ValueError`` naming the sheet and the header text sought when a
    required table is not found, and ``FileNotFoundError`` when ``path`` is absent.
    """
    resolved = Path(path)
    wb = openpyxl.load_workbook(resolved, read_only=True, data_only=True)
    try:
        schedule = _read_schedule(_get_sheet(wb, "SUMMARY (ER)"), "SUMMARY (ER)")
        decay_rates, doc_fractions = _read_decay_and_doc(
            _get_sheet(wb, "Waste Parameters"), "Waste Parameters"
        )
        composition, waste_by_year, biomethanization_fraction = _read_waste_projection(
            _get_sheet(wb, "Waste Projection"), "Waste Projection"
        )
        electricity, separated = _read_electricity_with_cross_check(wb)
    finally:
        wb.close()
    return RegisteredAcm0022Workbook(
        path=resolved,
        schedule=schedule,
        decay_rates=decay_rates,
        doc_fractions=doc_fractions,
        composition=composition,
        waste_by_year=waste_by_year,
        electricity_mwh_by_year=electricity,
        biomethanization_fraction=biomethanization_fraction,
        separated_baseline_by_year=separated,
    )


def to_project_input_patch(workbook: RegisteredAcm0022Workbook) -> dict[str, Any]:
    """Build the nested ProjectInput patch sourced from the workbook."""
    zone = infer_climate_zone({k: v.value for k, v in workbook.decay_rates.items()})
    technology: dict[str, Any] = {
        "waste_composition": [
            {
                "waste_type": waste_type,
                "mass_fraction": lineage.value,
                "source": (
                    f'registered ER workbook, sheet "Waste Projection", cell {lineage.cell}'
                ),
            }
            for waste_type, lineage in workbook.composition.items()
        ],
        "annual_waste_by_year": [v.value for v in workbook.waste_by_year],
        "energy_generation_mwh_by_year": [v.value for v in workbook.electricity_mwh_by_year],
        "biomethanization_suitable_fraction": (
            workbook.biomethanization_fraction.value
            if workbook.biomethanization_fraction is not None
            else None
        ),
    }
    patch: dict[str, Any] = {"location": {}, "technology": technology}
    if zone is not None:
        patch["location"] = {"climate_zone": zone}
    return patch
