"""Tests for the registered-workbook reader (PHASE-05)."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from pdd_agent.reconcile.workbook import (
    inegol_workbook_path,
    infer_climate_zone,
    read_acm0022_workbook,
    to_project_input_patch,
)


def _make_workbook(
    path: Path,
    *,
    years: int = 2,
    separated_override: list[float] | None = None,
    drop_waste_parameters: bool = False,
) -> Path:
    """Build a synthetic registered-style workbook in ``tmp_path``."""
    wb = openpyxl.Workbook()

    summary = wb.active
    summary.title = "SUMMARY (ER)"
    summary.append(
        [
            "Days",
            "Year",
            "Baseline Emissions (tCO2e)",
            "Project Emissions (tCO2e)",
            "Leak Emissions (tCO2e)",
            "Emission Reductions (tCO2e)",
        ]
    )
    be = [1000.0, 2000.0][:years]
    pe = [100.0, 200.0][:years]
    summary.append([1, "2020-12-31", 14, 16, 0, -2])  # stub row (Days < 300)
    labels = ["01/01/2021 - 12/31/2021", "01/01/2022 - 12/31/2022"][:years]
    for label, b, p in zip(labels, be, pe):
        summary.append([365, label, b, p, 0, b - p])
    summary.append(["TOTAL", "", sum(be), sum(pe), 0, sum(be) - sum(pe)])

    if not drop_waste_parameters:
        params = wb.create_sheet("Waste Parameters")
        params.append(["Waste", "DOCj", "kj"])
        params.append(["Wood and wood products", 43, 0.02])
        params.append(["Pulp, paper and cardboard (other than sludge)", 40, 0.04])
        params.append(["Food, food waste, beverages and tobacco", 15, 0.06])
        params.append(["Textiles", 24, 0.04])
        params.append(["Garden, yard and park waste", 20, 0.05])

    projection = wb.create_sheet("Waste Projection")
    projection.append(["Year", "Total Waste (A) (ton/year)", "Biomethanization (ton/year)"])
    totals = [10000.0, 20000.0][:years]
    bios = [4312.0, 8624.0][:years]
    for label, total, bio in zip(labels, totals, bios):
        projection.append([label, total, bio])
    projection.append([])
    projection.append(["Waste Type", "At the Integrated Solid Waste Facility"])
    projection.append(["Wood and wood products", 0.1])
    projection.append(["Food, food waste, beverages and tobacco", 0.6])
    projection.append(["Glass, plastic, metal, other inert waste", 0.3])

    project = wb.create_sheet("Project Emissions")
    project.append(["Year", "ECBL,k,y"])
    project.append(["2020-12-31", 0.0])
    for label, mwh in zip(labels, [100.0, 200.0][:years]):
        project.append([label, mwh])

    baseline = wb.create_sheet("Baseline Emissions (Total)")
    baseline.append(["Crediting Period", "Baseline Emissions from separate electricity generation"])
    baseline.append(["2020-12-31", 0.0])
    separated = (
        separated_override
        if separated_override is not None
        else [mwh * 0.541 * 1.09101 for mwh in [100.0, 200.0][:years]]
    )
    for label, value in zip(labels, separated):
        baseline.append([label, value])

    wb.save(path)
    return path


def test_schedule_reads_crediting_years_with_lineage(tmp_path: Path):
    wb = read_acm0022_workbook(_make_workbook(tmp_path / "synth.xlsx"))
    assert [v.value for v in wb.schedule.net] == [900.0, 1800.0]
    assert [v.value for v in wb.schedule.baseline] == [1000.0, 2000.0]
    assert [v.value for v in wb.schedule.leakage] == [0.0, 0.0]
    assert len(wb.schedule.year_labels) == 2  # the one-day stub row is absent
    assert wb.schedule.net[0].sheet == "SUMMARY (ER)"
    assert wb.schedule.net[0].cell == "F3"


def test_decay_and_doc_values(tmp_path: Path):
    wb = read_acm0022_workbook(_make_workbook(tmp_path / "synth.xlsx"))
    assert wb.decay_rates["food_waste"].value == pytest.approx(0.06)
    assert wb.doc_fractions["wood"].value == pytest.approx(0.43)


def test_infer_climate_zone_matches_all_supplied_values():
    assert (
        infer_climate_zone(
            {
                "wood": 0.02,
                "paper_cardboard": 0.04,
                "food_waste": 0.06,
                "textiles": 0.04,
                "garden_waste": 0.05,
            }
        )
        == "boreal_temperate_dry"
    )
    assert infer_climate_zone({"food_waste": 0.185}) == "boreal_temperate_wet"
    assert infer_climate_zone({"food_waste": 0.123}) is None


def test_biomethanization_fraction(tmp_path: Path):
    wb = read_acm0022_workbook(_make_workbook(tmp_path / "synth.xlsx"))
    assert wb.biomethanization_fraction is not None
    assert wb.biomethanization_fraction.value == pytest.approx(0.4312)


def test_ecbl_cross_check_failure_names_ecbl(tmp_path: Path):
    with pytest.raises(ValueError, match="ECBL"):
        read_acm0022_workbook(
            _make_workbook(tmp_path / "synth.xlsx", separated_override=[999, 999])
        )


def test_missing_waste_parameters_sheet_names_sheet(tmp_path: Path):
    with pytest.raises(ValueError, match="Waste Parameters"):
        read_acm0022_workbook(_make_workbook(tmp_path / "synth.xlsx", drop_waste_parameters=True))


def test_to_project_input_patch_shape(tmp_path: Path):
    wb = read_acm0022_workbook(_make_workbook(tmp_path / "synth.xlsx"))
    patch = to_project_input_patch(wb)
    entries = patch["technology"]["waste_composition"]
    assert all(set(e) == {"waste_type", "mass_fraction", "source"} for e in entries)
    inert = next(e for e in entries if e["waste_type"] == "inert")
    assert inert["mass_fraction"] == pytest.approx(0.3)
    assert patch["technology"]["annual_waste_by_year"] == [10000.0, 20000.0]
    assert patch["technology"]["energy_generation_mwh_by_year"] == [100.0, 200.0]
    assert patch["location"]["climate_zone"] == "boreal_temperate_dry"


@pytest.mark.corpus
def test_real_inegol_workbook_totals_and_zone():
    path = inegol_workbook_path()
    if not path.exists():
        pytest.skip(f"registered workbook absent: {path}")
    wb = read_acm0022_workbook(path)
    assert sum(v.value for v in wb.schedule.net) == 730000
    assert len(wb.schedule.net) == 7
    assert (
        infer_climate_zone({k: v.value for k, v in wb.decay_rates.items()})
        == "boreal_temperate_dry"
    )
